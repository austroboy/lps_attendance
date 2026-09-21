"""
The blank workbook people actually fill in.

It carries three things beyond a header row: dropdowns wired to whatever
classes, shifts, versions and groups already exist, a phone column forced to
text so 01712345678 does not become 1712345678, and a sheet of instructions.
Those three details are the difference between an import that works first time
and an afternoon of "why are all the numbers wrong".
"""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Group, SchoolClass, Section, Shift, Student, Version

HEADERS = [
    ("admission_no", 16, "Required. Unique. Re-uploading the same number updates that student."),
    ("roll_no", 9, "Optional."),
    ("full_name", 26, "Required."),
    ("class", 16, "Required. Created if it does not exist."),
    ("section", 11, "Required. Created if it does not exist."),
    ("shift", 12, "Optional. Morning, Day…"),
    ("version", 12, "Optional. Bangla, English."),
    ("group", 18, "Optional. Science, Business Studies… Leave blank below class nine."),
    ("guardian_name", 22, "Optional."),
    ("guardian_phone", 16, "Where SMS goes. 01XXXXXXXXX."),
    ("student_phone", 16, "Optional."),
    ("device_user_id", 15, "The number enrolled on the terminal. Fill in later if unknown."),
    ("is_active", 10, "yes or no. Blank means yes."),
]

HEAD_FILL = PatternFill("solid", fgColor="12454F")
REQ_FILL = PatternFill("solid", fgColor="C2703D")
NOTE_FILL = PatternFill("solid", fgColor="EEF3F3")
REQUIRED = {"admission_no", "full_name", "class", "section"}


def _sample_rows():
    """Prefer the school's own data so the example looks familiar."""
    rows = []
    for student in (Student.objects
                    .select_related("section", "section__school_class", "section__shift",
                                    "section__version", "section__group")
                    .filter(is_active=True)[:3]):
        section = student.section
        rows.append([
            student.admission_no, student.roll_no, student.full_name,
            section.school_class.name, section.name,
            section.shift.name if section.shift_id else "",
            section.version.name if section.version_id else "",
            section.group.name if section.group_id else "",
            student.guardian_name, student.guardian_phone, student.student_phone,
            student.device_user_id, "yes" if student.is_active else "no",
        ])
    if rows:
        return rows
    return [
        ["2026-0001", "1", "Rahim Uddin", "Class Six", "A", "Morning", "Bangla", "",
         "Karim Uddin", "01712345678", "", "101", "yes"],
        ["2026-0002", "2", "Fatima Akter", "Class Six", "A", "Morning", "Bangla", "",
         "Nasir Ahmed", "01812345678", "", "102", "yes"],
        ["2026-0003", "1", "Sabbir Hossain", "Class Nine", "A", "Day", "English", "Science",
         "Jamal Hossain", "01912345678", "", "103", "yes"],
    ]


def build_sample_workbook() -> bytes:
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Students"
    sheet.freeze_panes = "A2"

    # Hints live in cell notes on the header, not in a second row of the sheet.
    # A hint row looks helpful and then gets imported as a student: the first
    # version of this file created a class called "Required. Created if it does
    # not exist." Only row 1 is not data, and only row 1 exists.
    for index, (name, width, note) in enumerate(HEADERS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width

        head = sheet.cell(row=1, column=index, value=name)
        head.font = Font(bold=True, color="FFFFFF", size=11)
        head.fill = REQ_FILL if name in REQUIRED else HEAD_FILL
        head.alignment = Alignment(horizontal="center", vertical="center")
        comment = Comment(note, "LPS Attendance")
        comment.width, comment.height = 260, 90
        head.comment = comment

    sheet.row_dimensions[1].height = 22

    for offset, row in enumerate(_sample_rows(), start=2):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=offset, column=column, value=value)
            if HEADERS[column - 1][0] in ("guardian_phone", "student_phone",
                                          "admission_no", "device_user_id", "roll_no"):
                cell.number_format = "@"

    # Force the identifier and phone columns to text. This is a *column* style,
    # not a per-cell one: styling six thousand cells individually would create
    # six thousand empty rows and a 77 KB file that Excel reports as already
    # full. A column style costs nothing and applies to every row the person
    # ever types or pastes into, however many that is.
    for name in ("admission_no", "roll_no", "guardian_phone", "student_phone",
                 "device_user_id"):
        column = [h[0] for h in HEADERS].index(name) + 1
        dimension = sheet.column_dimensions[get_column_letter(column)]
        dimension.number_format = "@"
        dimension.customFormat = True

    _add_dropdowns(sheet)
    _add_reference_sheet(workbook)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _dropdown(sheet, column_name, values, prompt):
    if not values:
        return
    joined = ",".join(v.replace(",", " ") for v in values)
    if len(joined) > 240:   # Excel's inline list limit; skip rather than corrupt
        return
    column = [h[0] for h in HEADERS].index(column_name) + 1
    letter = get_column_letter(column)
    rule = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True,
                          showDropDown=False)
    rule.prompt = prompt
    rule.promptTitle = column_name
    sheet.add_data_validation(rule)
    rule.add(f"{letter}2:{letter}5000")


def _add_dropdowns(sheet):
    _dropdown(sheet, "class", list(SchoolClass.objects.values_list("name", flat=True)),
              "Pick one, or type a new class name to create it.")
    _dropdown(sheet, "shift", list(Shift.objects.values_list("name", flat=True)),
              "Pick one, or type a new shift to create it.")
    _dropdown(sheet, "version", list(Version.objects.values_list("name", flat=True)),
              "Pick one, or type a new version to create it.")
    _dropdown(sheet, "group", list(Group.objects.values_list("name", flat=True)),
              "Pick one, or leave blank below class nine.")
    _dropdown(sheet, "is_active", ["yes", "no"], "Blank means yes.")


def _add_reference_sheet(workbook):
    sheet = workbook.create_sheet("How to use this")
    sheet.column_dimensions["A"].width = 100

    lines = [
        ("How to fill this in", True),
        ("", False),
        ("Type your students from row 2 down. Do not delete or rename row 1 — the import "
         "reads the column names from it. Hover any column heading for a note about what "
         "goes in it.", False),
        ("", False),
        ("Orange columns are required: admission_no, full_name, class, section.", False),
        ("Everything else can be left blank.", False),
        ("", False),
        ("admission_no is the key. Upload the same file twice and nobody is duplicated — "
         "the second run just updates them. That means you can fix a typo in the "
         "spreadsheet and re-upload rather than editing students one by one.", False),
        ("", False),
        ("class, section, shift, version and group are created automatically if they do "
         "not exist yet. You do not need to set anything up first. The import tells you "
         "afterwards exactly what it created, so a typo like 'Clas Six' is easy to spot "
         "before it spreads.", False),
        ("", False),
        ("Phone numbers must keep their leading zero. This file already forces those "
         "columns to text. If you paste from elsewhere and see 1712345678, format the "
         "column as Text and retype the zero.", False),
        ("", False),
        ("device_user_id is the number the face terminal knows a student by. Leave it "
         "blank for now — you can fill it in later from Devices, Device users, once the "
         "terminals are enrolled.", False),
        ("", False),
        ("Tick 'Check the file first' when you upload. It reports every problem and "
         "everything it would create, without writing a single row.", False),
        ("", False),
        ("Current structure", True),
    ]

    row = 1
    for text, is_heading in lines:
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, size=13) if is_heading else Font(size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if text and not is_heading and len(text) > 90:
            sheet.row_dimensions[row].height = 30
        row += 1

    for label, queryset in (("Classes", SchoolClass.objects.values_list("name", flat=True)),
                            ("Shifts", Shift.objects.values_list("name", flat=True)),
                            ("Versions", Version.objects.values_list("name", flat=True)),
                            ("Groups", Group.objects.values_list("name", flat=True))):
        values = list(queryset)
        cell = sheet.cell(row=row, column=1,
                          value=f"{label}: " + (", ".join(values) if values else "none yet"))
        cell.font = Font(size=10)
        row += 1

    sections = Section.objects.select_related("school_class")[:40]
    if sections:
        row += 1
        sheet.cell(row=row, column=1, value="Sections that already exist:").font = Font(bold=True)
        row += 1
        for section in sections:
            sheet.cell(row=row, column=1, value=f"   {section}").font = Font(size=10)
            row += 1
