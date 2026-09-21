"""Tests for the bulk student import and the structures it creates."""
import io

from django.core.files.base import ContentFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from .import_models import ImportJob, ImportStatus
from .imports import run_job
from .models import ClassAccess, Group, SchoolClass, Section, Shift, Student, Teacher, Version
from .sample import build_sample_workbook

HEADER = ["admission_no", "roll_no", "full_name", "class", "section", "shift",
          "version", "group", "guardian_name", "guardian_phone", "student_phone",
          "device_user_id", "is_active"]


def workbook_bytes(rows, header=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Students"
    sheet.append(header or HEADER)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def run(content, name="roll.xlsx", **kwargs):
    job = ImportJob(original_name=name, **kwargs)
    job.upload.save(name, ContentFile(content), save=False)
    job.save()
    run_job(job)
    job.refresh_from_db()
    return job


def student_row(n, **over):
    row = {"admission_no": f"A{n:04d}", "roll_no": str(n), "full_name": f"Student {n}",
           "class": "Class Nine", "section": "A", "shift": "Morning",
           "version": "Bangla", "group": "Science", "guardian_name": f"Guardian {n}",
           "guardian_phone": "01712345678", "student_phone": "",
           "device_user_id": str(1000 + n), "is_active": "yes"}
    row.update(over)
    return [row[key] for key in HEADER]


class StructureCreationTests(TestCase):
    def test_everything_is_created_from_the_file(self):
        job = run(workbook_bytes([student_row(1), student_row(2)]))
        self.assertEqual(job.status, ImportStatus.DONE)
        self.assertEqual(job.created, 2)
        self.assertEqual(SchoolClass.objects.count(), 1)
        self.assertEqual(Shift.objects.count(), 1)
        self.assertEqual(Version.objects.count(), 1)
        self.assertEqual(Group.objects.count(), 1)
        self.assertEqual(Section.objects.count(), 1)
        self.assertEqual(job.new_classes, ["Class Nine"])

    def test_same_section_letter_splits_by_shift_and_version(self):
        run(workbook_bytes([
            student_row(1, shift="Morning", version="Bangla"),
            student_row(2, shift="Day", version="Bangla"),
            student_row(3, shift="Morning", version="English"),
        ]))
        # Same class, same letter A — but three genuinely different rooms.
        self.assertEqual(Section.objects.count(), 3)

    def test_group_is_optional(self):
        run(workbook_bytes([student_row(1, **{"class": "Class Six", "group": ""})]))
        section = Section.objects.get()
        self.assertIsNone(section.group_id)

    def test_classes_sort_in_school_order_not_alphabetically(self):
        run(workbook_bytes([
            student_row(1, **{"class": "Class Ten"}),
            student_row(2, **{"class": "Class Six"}),
            student_row(3, **{"class": "Class Nine"}),
        ]))
        names = list(SchoolClass.objects.values_list("name", flat=True))
        self.assertEqual(names, ["Class Six", "Class Nine", "Class Ten"])

    def test_existing_structure_is_reused_not_duplicated(self):
        run(workbook_bytes([student_row(1)]))
        run(workbook_bytes([student_row(2)]), name="second.xlsx")
        self.assertEqual(SchoolClass.objects.count(), 1)
        self.assertEqual(Section.objects.count(), 1)
        self.assertEqual(Student.objects.count(), 2)


class RowHandlingTests(TestCase):
    def test_reupload_updates_and_never_duplicates(self):
        run(workbook_bytes([student_row(1), student_row(2)]))
        job = run(workbook_bytes([student_row(1, full_name="Renamed"), student_row(2)]),
                  name="again.xlsx")
        self.assertEqual(job.created, 0)
        self.assertEqual(job.updated, 2)
        self.assertEqual(Student.objects.count(), 2)
        self.assertEqual(Student.objects.get(admission_no="A0001").full_name, "Renamed")

    def test_a_student_who_changed_class_is_moved(self):
        run(workbook_bytes([student_row(1, **{"class": "Class Nine"})]))
        run(workbook_bytes([student_row(1, **{"class": "Class Ten"})]), name="promoted.xlsx")
        student = Student.objects.get(admission_no="A0001")
        self.assertEqual(student.section.school_class.name, "Class Ten")
        self.assertEqual(Student.objects.count(), 1)

    def test_bad_rows_are_skipped_and_explained(self):
        job = run(workbook_bytes([
            student_row(1),
            student_row(2, admission_no=""),
            student_row(3, full_name=""),
            student_row(4, **{"class": ""}),
            student_row(1, full_name="Duplicate"),
        ]))
        self.assertEqual(job.created, 1)
        self.assertEqual(job.skipped, 4)
        self.assertEqual(len(job.problems), 4)
        self.assertTrue(any("more than once" in p for p in job.problems))

    def test_phone_numbers_keep_their_leading_zero(self):
        run(workbook_bytes([
            student_row(1, guardian_phone="1712345678"),      # Excel ate the zero
            student_row(2, guardian_phone=1712345678),        # stored as a number
            student_row(3, guardian_phone="+880 1712-345678"),
        ]))
        numbers = set(Student.objects.values_list("guardian_phone", flat=True))
        self.assertEqual(numbers, {"01712345678", "8801712345678"})

    def test_whitespace_is_tidied(self):
        run(workbook_bytes([student_row(1, full_name="  Rahim   Uddin  ")]))
        self.assertEqual(Student.objects.get().full_name, "Rahim Uddin")

    def test_a_pasted_note_row_cannot_invent_a_class(self):
        job = run(workbook_bytes([
            ["Required. Unique. Re-uploading updates that student.", "", "Required.",
             "Required. Created if it does not exist.", "Required.", "", "", "", "", "", "", "", ""],
            student_row(1),
        ]))
        self.assertEqual(job.created, 1)
        self.assertEqual(job.skipped, 1)
        self.assertFalse(SchoolClass.objects.filter(name__icontains="Required").exists())

    def test_deactivate_missing_only_touches_sections_in_the_file(self):
        run(workbook_bytes([student_row(1), student_row(2)]))
        other = Section.objects.create(
            school_class=SchoolClass.objects.create(name="Class Two", order=4), name="A")
        outsider = Student.objects.create(admission_no="Z1", full_name="Elsewhere",
                                          section=other)
        job = run(workbook_bytes([student_row(1)]), name="trimmed.xlsx",
                  deactivate_missing=True)
        self.assertEqual(job.deactivated, 1)
        self.assertFalse(Student.objects.get(admission_no="A0002").is_active)
        outsider.refresh_from_db()
        self.assertTrue(outsider.is_active)


class DryRunTests(TestCase):
    def test_dry_run_writes_nothing(self):
        job = run(workbook_bytes([student_row(1), student_row(2)]), dry_run=True)
        self.assertEqual(job.status, ImportStatus.DONE)
        self.assertEqual(Student.objects.count(), 0)
        self.assertEqual(SchoolClass.objects.count(), 0)
        self.assertEqual(Section.objects.count(), 0)

    def test_dry_run_predicts_the_real_run(self):
        rows = [student_row(n, shift="Morning" if n % 2 else "Day") for n in range(1, 21)]
        preview = run(workbook_bytes(rows), name="preview.xlsx", dry_run=True)
        real = run(workbook_bytes(rows), name="real.xlsx")

        self.assertEqual(preview.created, real.created)
        self.assertEqual(preview.updated, real.updated)
        self.assertEqual(sorted(preview.new_sections), sorted(real.new_sections))
        self.assertEqual(sorted(preview.new_classes), sorted(real.new_classes))

    def test_dry_run_separates_new_from_existing(self):
        run(workbook_bytes([student_row(1)]))
        job = run(workbook_bytes([student_row(1), student_row(2)]),
                  name="mixed.xlsx", dry_run=True)
        self.assertEqual(job.created, 1)
        self.assertEqual(job.updated, 1)


class FileFormatTests(TestCase):
    def test_csv_with_friendly_column_names(self):
        csv = (b"Admission No,Name,Grade,Sec,Mobile\n"
               b"C-1,CSV Student,Class Six,B,01812345678\n")
        job = run(csv, name="roll.csv")
        self.assertEqual(job.created, 1)
        student = Student.objects.get(admission_no="C-1")
        self.assertEqual(student.full_name, "CSV Student")
        self.assertEqual(student.guardian_phone, "01812345678")

    def test_missing_required_columns_fails_with_a_useful_message(self):
        job = run(workbook_bytes([["x", "y"]], header=["foo", "bar"]))
        self.assertEqual(job.status, ImportStatus.FAILED)
        self.assertIn("admission_no", job.error)

    def test_a_file_that_is_not_a_workbook_says_so_in_plain_words(self):
        job = run(b"this is not a spreadsheet")
        self.assertEqual(job.status, ImportStatus.FAILED)
        self.assertIn("Save as", job.error)

    def test_blank_rows_are_ignored(self):
        content = workbook_bytes([student_row(1), [None] * len(HEADER), student_row(2)])
        job = run(content)
        self.assertEqual(job.total_rows, 2)
        self.assertEqual(job.created, 2)


class SampleWorkbookTests(TestCase):
    def test_sample_has_one_header_row_and_no_phantom_rows(self):
        book = load_workbook(io.BytesIO(build_sample_workbook()))
        sheet = book["Students"]
        self.assertEqual([c.value for c in sheet[1]], HEADER)
        # Header plus the example rows, nothing more. Per-cell styling used to
        # leave thousands of empty rows behind.
        self.assertLessEqual(sheet.max_row, 5)

    def test_sample_keeps_identifier_columns_as_text(self):
        book = load_workbook(io.BytesIO(build_sample_workbook()))
        sheet = book["Students"]
        self.assertEqual(sheet.column_dimensions["J"].number_format, "@")
        self.assertEqual(sheet.column_dimensions["A"].number_format, "@")

    def test_sample_round_trips_straight_back_in(self):
        book = load_workbook(io.BytesIO(build_sample_workbook()))
        sheet = book["Students"]
        for n in range(1, 11):
            sheet.append([f"RT-{n}", str(n), f"Round Trip {n}", "Class Twelve", "C",
                          "Day", "English", "Science", "Guardian", "01722000000",
                          "", str(9000 + n), "yes"])
        buffer = io.BytesIO()
        book.save(buffer)

        job = run(buffer.getvalue(), name="filled.xlsx")
        self.assertEqual(job.status, ImportStatus.DONE)
        self.assertEqual(job.skipped, 0)
        self.assertTrue(Student.objects.filter(admission_no="RT-5").exists())
        self.assertFalse(SchoolClass.objects.filter(name__icontains="Required").exists())

    def test_sample_carries_an_instructions_sheet(self):
        book = load_workbook(io.BytesIO(build_sample_workbook()))
        self.assertIn("How to use this", book.sheetnames)


class ScaleTests(TestCase):
    def test_two_thousand_rows_use_a_handful_of_queries(self):
        """
        The whole point of the two-pass design.

        Row-by-row update_or_create would be roughly 10,000 queries and minutes
        of wall clock on a small VPS. Caching the lookups and using bulk_create
        keeps it near a hundred.
        """
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        rows = [student_row(n, **{"class": f"Class {n % 5}", "section": "ABCD"[n % 4]})
                for n in range(1, 2001)]
        job = ImportJob(original_name="big.xlsx")
        job.upload.save("big.xlsx", ContentFile(workbook_bytes(rows)), save=False)
        job.save()

        with CaptureQueriesContext(connection) as captured:
            run_job(job)

        job.refresh_from_db()
        self.assertEqual(job.created, 2000)
        self.assertEqual(Student.objects.count(), 2000)
        self.assertLess(len(captured), 300,
                        f"{len(captured)} queries for 2,000 rows — the lookup cache or the "
                        f"bulk path has regressed to per-row queries")


class EnvFileTests(TestCase):
    """
    .env has to actually be read.

    Every setting comes from the environment and .env.example says to copy it to
    .env — but for several versions nothing loaded the file, so a gateway token
    or broker URL put there was silently ignored.
    """

    def test_values_are_read_and_real_env_vars_still_win(self):
        import os
        import tempfile
        from pathlib import Path

        from config.settings import _load_env_file

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / ".env"
            path.write_text(
                "# a comment\n"
                "\n"
                'QUOTED_VALUE="hello"\n'
                "PLAIN_VALUE = spaced \n"
                "ALREADY_SET=from_file\n"
                "not a pair\n"
            )
            os.environ["ALREADY_SET"] = "from_shell"
            try:
                _load_env_file(path)
                self.assertEqual(os.environ["QUOTED_VALUE"], "hello")
                self.assertEqual(os.environ["PLAIN_VALUE"], "spaced")
                # A real environment variable beats the file.
                self.assertEqual(os.environ["ALREADY_SET"], "from_shell")
            finally:
                for key in ("QUOTED_VALUE", "PLAIN_VALUE", "ALREADY_SET"):
                    os.environ.pop(key, None)

    def test_a_missing_file_is_not_an_error(self):
        from pathlib import Path

        from config.settings import _load_env_file

        _load_env_file(Path("/nowhere/at/all/.env"))


class RunWithoutCeleryTests(TestCase):
    """
    The import must be usable with no broker at all.

    A laptop without Redis, or a server whose worker has stopped, still needs a
    way to get the students in. The old Run now button only re-queued to the
    same unreachable broker, so it appeared to do nothing.
    """

    def setUp(self):
        from accounts.models import Role, User

        self.admin = User.objects.create(username="boss", role=Role.SUPERADMIN,
                                         is_staff=True, is_superuser=True)
        self.admin.set_password("x")
        self.admin.save()
        self.client.force_login(self.admin)

        self.job = ImportJob(original_name="roll.xlsx")
        self.job.upload.save("roll.xlsx",
                             ContentFile(workbook_bytes([student_row(n) for n in range(1, 26)])),
                             save=False)
        self.job.save()

    def test_run_now_imports_without_a_worker(self):
        response = self.client.get(f"/academics/students/import/{self.job.pk}/run/")
        self.assertEqual(response.status_code, 302)

        self.job.refresh_from_db()
        self.assertEqual(self.job.status, ImportStatus.DONE)
        self.assertEqual(self.job.created, 25)
        self.assertEqual(Student.objects.count(), 25)

    def test_run_now_reports_a_bad_file_instead_of_raising(self):
        job = ImportJob(original_name="junk.xlsx")
        job.upload.save("junk.xlsx", ContentFile(b"not a workbook"), save=False)
        job.save()

        response = self.client.get(f"/academics/students/import/{job.pk}/run/")
        self.assertEqual(response.status_code, 302)
        job.refresh_from_db()
        self.assertEqual(job.status, ImportStatus.FAILED)
        self.assertTrue(job.error)

    def test_running_twice_does_not_double_import(self):
        self.client.get(f"/academics/students/import/{self.job.pk}/run/")
        self.client.get(f"/academics/students/import/{self.job.pk}/run/")
        self.assertEqual(Student.objects.count(), 25)

    def test_a_teacher_cannot_trigger_an_import(self):
        from accounts.models import Role, User

        teacher = User.objects.create(username="teach", role=Role.TEACHER)
        teacher.set_password("x")
        teacher.save()
        self.client.force_login(teacher)

        response = self.client.get(f"/academics/students/import/{self.job.pk}/run/")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(Student.objects.count(), 0)


class RosterTests(TestCase):
    """Card numbers, face enrolment, and taking a student off the roll."""

    def setUp(self):
        import json
        import struct

        from accounts.models import Role, User
        from devices.models import Device

        self.frame = lambda p: (
            struct.pack("<I", len(json.dumps(p, separators=(",", ":")).encode() + b"\x00"))
            + json.dumps(p, separators=(",", ":")).encode() + b"\x00")

        self.device = Device.objects.create(serial_number="SIM001", name="Gate",
                                            purpose="STUDENT")
        six = SchoolClass.objects.create(name="Class Six", order=8)
        eight = SchoolClass.objects.create(name="Class Eight", order=10)
        self.mine = Section.objects.create(school_class=six, name="A")
        self.theirs = Section.objects.create(school_class=eight, name="A")

        self.student = Student.objects.create(admission_no="S1", full_name="Rahim",
                                              section=self.mine)
        self.outsider = Student.objects.create(admission_no="S2", full_name="Karim",
                                               section=self.theirs)

        self.admin = User.objects.create(username="boss", role=Role.SUPERADMIN,
                                         is_staff=True, is_superuser=True)
        self.admin.set_password("x")
        self.admin.save()

        teacher_user = User.objects.create(username="teach", role=Role.TEACHER)
        teacher_user.set_password("x")
        teacher_user.save()
        teacher = Teacher.objects.create(user=teacher_user, employee_code="T-9",
                                         full_name="Nasrin")
        ClassAccess.objects.create(teacher=teacher, school_class=six)
        self.teacher_user = teacher_user

    # -- cards ----------------------------------------------------------
    def test_card_is_cleaned_and_saved(self):
        self.client.force_login(self.admin)
        response = self.client.post(f"/academics/students/{self.student.pk}/rfid/",
                                    {"rfid": " 00-12 ab34 "})
        self.assertEqual(response.status_code, 200)
        self.student.refresh_from_db()
        self.assertEqual(self.student.rfid_number, "0012AB34")

    def test_card_is_pushed_to_the_terminals(self):
        from devices.models import DeviceCommand

        self.client.force_login(self.admin)
        self.client.post(f"/academics/students/{self.student.pk}/rfid/", {"rfid": "ABC1"})
        command = DeviceCommand.objects.get(cmd_code="SET_USER_INFO")
        self.assertEqual(command.cmd_param["card_number"], "ABC1")

    def test_the_same_card_cannot_go_on_two_students(self):
        self.client.force_login(self.admin)
        self.client.post(f"/academics/students/{self.student.pk}/rfid/", {"rfid": "DUP1"})
        response = self.client.post(f"/academics/students/{self.outsider.pk}/rfid/",
                                    {"rfid": "DUP1"})
        self.assertEqual(response.status_code, 409)
        self.assertIn("already on", response.json()["error"])
        self.outsider.refresh_from_db()
        self.assertEqual(self.outsider.rfid_number, "")

    # -- face -----------------------------------------------------------
    def test_starting_enrolment_assigns_a_device_id(self):
        self.assertEqual(self.student.device_user_id, "")
        self.client.force_login(self.admin)
        response = self.client.post(f"/academics/students/{self.student.pk}/face/",
                                    {"action": "start"})
        data = response.json()
        self.assertEqual(data["face"], "PENDING")
        self.assertTrue(data["device_user_id"])
        self.student.refresh_from_db()
        self.assertEqual(self.student.face_status, "PENDING")

    def test_device_ids_do_not_collide(self):
        from .roster import next_device_user_id

        Student.objects.filter(pk=self.outsider.pk).update(device_user_id="7")
        Teacher.objects.filter(employee_code="T-9").update(device_user_id="8")
        self.assertEqual(next_device_user_id(), "9")

    def test_terminal_enrolment_ticks_the_row(self):
        self.client.force_login(self.admin)
        data = self.client.post(f"/academics/students/{self.student.pk}/face/",
                                {"action": "start"}).json()

        self.client.post("/ebkn/", data=self.frame({
            "user_id": data["device_user_id"], "user_name": "Rahim",
            "enroll_data_array": [{"backup_number": 10}],
        }), content_type="application/octet-stream",
            headers={"dev_id": "SIM001", "request_code": "realtime_enroll_data"})

        self.student.refresh_from_db()
        self.assertEqual(self.student.face_status, "DONE")
        self.assertIn("face", self.student.face_detail)
        self.assertEqual(self.student.face_device, self.device)

    def test_a_face_punch_alone_counts_as_enrolled(self):
        """Firmware that never pushes enrolment data still gets us there."""
        from django.utils import timezone

        Student.objects.filter(pk=self.student.pk).update(device_user_id="55")
        stamp = timezone.localdate().strftime("%Y%m%d")
        self.client.post("/ebkn/", data=self.frame({
            "user_id": "55", "verify_mode": "FACE", "io_mode": 1,
            "io_time": f"{stamp}074500",
        }), content_type="application/octet-stream",
            headers={"dev_id": "SIM001", "request_code": "realtime_glog"})

        self.student.refresh_from_db()
        self.assertEqual(self.student.face_status, "DONE")

    def test_enrolment_can_be_cleared_and_redone(self):
        self.client.force_login(self.admin)
        self.student.mark_face_enrolled(device=self.device, detail="face")
        self.client.post(f"/academics/students/{self.student.pk}/face/", {"action": "reset"})
        self.student.refresh_from_db()
        self.assertEqual(self.student.face_status, "NONE")
        self.assertIsNone(self.student.face_enrolled_at)

    # -- leaving the roll ------------------------------------------------
    def test_removal_keeps_history_and_clears_the_terminals(self):
        from devices.models import DeviceCommand

        Student.objects.filter(pk=self.student.pk).update(device_user_id="42")
        self.client.force_login(self.admin)
        self.client.post(f"/academics/students/{self.student.pk}/active/",
                         {"active": "0", "reason": "TRANSFERRED", "note": "Moved"})

        self.student.refresh_from_db()
        self.assertFalse(self.student.is_active)
        self.assertEqual(self.student.leave_reason, "TRANSFERRED")
        self.assertIsNotNone(self.student.left_on)
        self.assertEqual(self.student.deactivated_by, self.admin)
        # Still in the database — last month's reports must not change.
        self.assertTrue(Student.objects.filter(pk=self.student.pk).exists())

        command = DeviceCommand.objects.get(cmd_code="DELETE_USER")
        self.assertEqual(command.cmd_param["user_id"], "42")

    def test_removal_can_leave_them_on_the_terminals(self):
        from devices.models import DeviceCommand

        Student.objects.filter(pk=self.student.pk).update(device_user_id="42")
        self.client.force_login(self.admin)
        self.client.post(f"/academics/students/{self.student.pk}/active/",
                         {"active": "0", "keep_on_devices": "1"})
        self.assertFalse(DeviceCommand.objects.filter(cmd_code="DELETE_USER").exists())

    def test_restoring_puts_them_back_and_expects_a_re_enrolment(self):
        Student.objects.filter(pk=self.student.pk).update(device_user_id="42")
        self.student.refresh_from_db()
        self.student.mark_face_enrolled(device=self.device, detail="face")

        self.client.force_login(self.admin)
        self.client.post(f"/academics/students/{self.student.pk}/active/", {"active": "0"})
        self.client.post(f"/academics/students/{self.student.pk}/active/", {"active": "1"})

        self.student.refresh_from_db()
        self.assertTrue(self.student.is_active)
        # The template lived on the terminal and was deleted with the user, so
        # claiming they still have a face would be a lie.
        self.assertEqual(self.student.face_status, "NONE")

    # -- who is allowed --------------------------------------------------
    def test_a_teacher_can_work_on_their_own_class(self):
        self.client.force_login(self.teacher_user)
        for url, body in (
            (f"/academics/students/{self.student.pk}/rfid/", {"rfid": "T1"}),
            (f"/academics/students/{self.student.pk}/face/", {"action": "start"}),
            (f"/academics/students/{self.student.pk}/active/", {"active": "0"}),
        ):
            self.assertEqual(self.client.post(url, body).status_code, 200, url)

    def test_a_teacher_cannot_touch_another_class(self):
        self.client.force_login(self.teacher_user)
        for url, body in (
            (f"/academics/students/{self.outsider.pk}/rfid/", {"rfid": "T2"}),
            (f"/academics/students/{self.outsider.pk}/face/", {"action": "start"}),
            (f"/academics/students/{self.outsider.pk}/active/", {"active": "0"}),
        ):
            self.assertEqual(self.client.post(url, body).status_code, 403, url)

        self.outsider.refresh_from_db()
        self.assertEqual(self.outsider.rfid_number, "")
        self.assertTrue(self.outsider.is_active)

    def test_a_student_account_can_do_none_of_it(self):
        from accounts.models import Role, User

        pupil = User.objects.create(username="pupil", role=Role.STUDENT)
        pupil.set_password("x")
        pupil.save()
        self.client.force_login(pupil)

        for path in ("rfid", "face", "active"):
            response = self.client.post(f"/academics/students/{self.student.pk}/{path}/", {})
            self.assertEqual(response.status_code, 403)

    def test_polling_is_filtered_to_the_teachers_own_class(self):
        self.client.force_login(self.teacher_user)
        response = self.client.get(
            f"/academics/students/states/?ids={self.student.pk},{self.outsider.pk}")
        returned = [row["id"] for row in response.json()["students"]]
        self.assertEqual(returned, [self.student.pk])


class RollNumberTests(TestCase):
    """At this school the admission ID is the roll number."""

    def setUp(self):
        six = SchoolClass.objects.create(name="Class Six", order=8)
        self.section = Section.objects.create(school_class=six, name="A")

    def test_blank_roll_takes_the_admission_number(self):
        student = Student.objects.create(admission_no="LPS-JS-260044", full_name="Ahnaf",
                                         section=self.section)
        self.assertEqual(student.roll_no, "LPS-JS-260044")

    def test_a_typed_roll_is_kept_whatever_its_shape(self):
        for roll in ("12", "A-07", "LPS BM 250927", "রোল-৫"):
            student = Student.objects.create(admission_no=f"X{roll}", full_name="S",
                                             section=self.section, roll_no=roll)
            student.refresh_from_db()
            self.assertEqual(student.roll_no, roll)

    def test_import_fills_a_blank_roll_too(self):
        # bulk_create skips save(), so the importer has to do it itself.
        run(workbook_bytes([student_row(1, roll_no=""), student_row(2, roll_no="R-2")]))
        self.assertEqual(Student.objects.get(admission_no="A0001").roll_no, "A0001")
        self.assertEqual(Student.objects.get(admission_no="A0002").roll_no, "R-2")

    def test_roster_does_not_repeat_the_id_under_the_name(self):
        from accounts.models import Role, User

        admin = User.objects.create(username="boss", role=Role.SUPERADMIN,
                                    is_staff=True, is_superuser=True)
        Student.objects.create(admission_no="LPS-JS-260044", full_name="Ahnaf",
                               section=self.section)
        self.client.force_login(admin)
        html = self.client.get("/academics/students/").content.decode()
        self.assertEqual(html.count("LPS-JS-260044"), 1)
